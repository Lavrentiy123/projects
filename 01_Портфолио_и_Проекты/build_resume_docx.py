# -*- coding: utf-8 -*-
"""
Генератор официального документа резюме Лаврентия Ямпурова в формате DOCX.
AI Product Manager / AI Solutions Architect / Lead Systems Analyst.
"""
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls


def create_resume():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    # Шапка
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    r_name = p_title.add_run("Лаврентий Ямпуров\n")
    r_name.bold = True
    r_name.font.size = Pt(18)
    r_name.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    r_role = p_title.add_run("AI Product Manager • AI Solutions Architect • Lead Systems Analyst")
    r_role.bold = True
    r_role.font.size = Pt(11.5)
    r_role.font.color.rgb = RGBColor(0x00, 0x72, 0xCE)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(6)
    p_meta.add_run("Контакты: +7 (981) 167-82-36 | ylv30072002@mail.ru | Telegram: @Lavr02 | Санкт-Петербург\n")
    p_meta.add_run("GitHub: github.com/Lavrentiy123 | Резюме актуально на 2026 год")

    def add_h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    # Профиль
    add_h("ЦЕЛЕВОЙ ПРОФИЛЬ И СИНЕРГИЯ КОМПЕТЕНЦИЙ")
    p_prof = doc.add_paragraph()
    p_prof.paragraph_format.space_after = Pt(4)
    p_prof.paragraph_format.line_spacing = 1.15
    p_prof.add_run(
        "Специалист редкого гибридного профиля: объединяю глубокое понимание архитектуры мультиагентных систем "
        "и прикладного машинного обучения с опытом управления цифровой трансформацией в тяжелой промышленности (ВГК) "
        "и академическим базисом в области стратегического менеджмента (Магистратура НИУ ИТМО по инноватике и стратегиям трансформации).\n"
        "Ключевой фокус: проектирование и сквозной запуск корпоративных ИИ-продуктов (Multi-Agent Systems, Advanced RAG, "
        "Predictive Maintenance, FastMCP, Decision Engines, On-Premise LLM) от выявления потребностей бизнеса и ТЗ "
        "до архитектуры, интеграции в ERP/DWH и защиты ROI перед C-level."
    )

    # Стек
    add_h("КЛЮЧЕВЫЕ НАВЫКИ И ТЕХНОЛОГИЧЕСКИЙ СТЕК")
    skills = [
        ("Мультиагентные системы и LLM: ", "Multi-Agent Orchestration, LangGraph (StateGraph), Pydantic v2, FastMCP (Anthropic Model Context Protocol), Evaluator-Optimizer (Reflexion Loop), Prompt Engineering, On-Premise vLLM / Ollama (Qwen, DeepSeek)."),
        ("Advanced RAG и Поиск: ", "Гибридный поиск (Dense bge-m3 + Sparse BM25 + RRF), кросс-энкодерный реранкинг (bge-reranker-large), Layout-парсеры (Docling, Marker), векторные базы Qdrant / Chroma, метрики качества Ragas."),
        ("Индустриальный ML и Big Data: ", "Python 3.9+, Scikit-learn, XGBoost, ClickHouse, Apache Spark streaming logic, MLflow, MLOps, анализ промышленной телеметрии (датчики MWD/LWD бурения, вибродиагностика насосов УЭЦН)."),
        ("Системный и Бизнес-анализ: ", "BPMN 2.0, UML, User Stories, Use Cases, BRD / FRD, проектирование REST API, JSON-RPC, интеграции с 1C:ERP."),
        ("Product & Delivery Management: ", "Agile, Scrum, Scrumban, приоритизация бэклога (RICE/ICE), юнит-экономика, CustDev, риск-менеджмент (Red Flags), стейкхолдер-менеджмент C-level."),
        ("BI и Отчетность: ", "Yandex DataLens, SQL, Excel/Power Query, программная генерация презентаций (python-pptx) и отчетов (python-docx).")
    ]
    for s_title, s_desc in skills:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.left_indent = Inches(0.15)
        p_s.paragraph_format.space_after = Pt(1.5)
        p_s.paragraph_format.line_spacing = 1.1
        r_st = p_s.add_run(f"• {s_title}")
        r_st.bold = True
        r_st.font.size = Pt(8.5)
        r_st.font.color.rgb = RGBColor(0x00, 0x4C, 0x87)
        r_sd = p_s.add_run(s_desc)
        r_sd.font.size = Pt(8.5)

    # Опыт работы
    add_h("ОПЫТ РАБОТЫ")

    p_w1 = doc.add_paragraph()
    p_w1.paragraph_format.space_before = Pt(3)
    p_w1.paragraph_format.space_after = Pt(1)
    r_w1_t = p_w1.add_run("Восточная горнорудная компания (ВГК) | Руководитель проектов цифровой трансформации")
    r_w1_t.bold = True
    r_w1_t.font.size = Pt(9.5)
    r_w1_t.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_w1_sub = doc.add_paragraph()
    p_w1_sub.paragraph_format.space_after = Pt(2)
    r_sub = p_w1_sub.add_run("Август 2026 — настоящее время | Горнорудная промышленность, добыча и обогащение полезных ископаемых")
    r_sub.italic = True
    r_sub.font.size = Pt(8)

    exp_points = [
        ("Интеллектуальный агент обработки дефектов ТОиР: ", "спроектировал и внедрил автоматизированный контур парсинга сообщений мастеров из мессенджера MAX в 1C:ERP через API с RAG-поиском по регламентам ремонтов. Результат: резкое сокращение времени регистрации инцидентов и исключение ошибок ручного ввода."),
        ("Контур раннего предупреждения отказов оборудования (ГТО): ", "архитектура предиктивного мониторинга карьерных экскаваторов и самосвалов по телеметрии вибро- и термодатчиков. Переход к предиктивному обслуживанию (PdM)."),
        ("Оценка эффективности буровзрывных работ (БВР): ", "автоматизация сопоставления параметров бурения и качества дробления породы для снижения себестоимости добычи.")
    ]
    for e_t, e_d in exp_points:
        p_e = doc.add_paragraph()
        p_e.paragraph_format.left_indent = Inches(0.15)
        p_e.paragraph_format.space_after = Pt(1.5)
        p_e.paragraph_format.line_spacing = 1.12
        r_et = p_e.add_run(f"– {e_t}")
        r_et.bold = True
        r_et.font.size = Pt(8.5)
        r_ed = p_e.add_run(e_d)
        r_ed.font.size = Pt(8.5)

    p_w2 = doc.add_paragraph()
    p_w2.paragraph_format.space_before = Pt(4)
    p_w2.paragraph_format.space_after = Pt(1)
    r_w2_t = p_w2.add_run("Руководитель прикладных ИИ-разработок и продуктовых решений | 2025 — настоящее время")
    r_w2_t.bold = True
    r_w2_t.font.size = Pt(9.5)
    r_w2_t.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_w2_sub = doc.add_paragraph()
    p_w2_sub.paragraph_format.space_after = Pt(2)
    r_sub2 = p_w2_sub.add_run("Личные end-to-end проекты и разработка в партнерстве со специалистами из технологических компаний и отраслевых подразделений")
    r_sub2.italic = True
    r_sub2.font.size = Pt(8)

    rd_projects = [
        ("Мультиагентный конвейер ДСИиУР (Газпром нефть): ", "сквозная система из 3 агентов на LangGraph: сбор отраслевых сигналов (Early Warnings), генерация инициатив с циклом рефлексии (Reflexion) по критериальной матрице, автоматическая сборка презентаций PPTX по брендбуку. Готовый прототип run.py, питч-дек на 21 слайд."),
        ("Multi-Agent Judge & Reflexion System: ", "архитектурный контур из 4-х агентов по паттерну Anthropic Building Effective Agents (Evaluator-Optimizer) с метриками Ragas для гарантированного отсечения галлюцинаций в корпоративных отчетах."),
        ("OilGas FastMCP Server: ", "сервер контекста по стандарту Model Context Protocol (FastMCP + Qdrant) для семантического поиска по нормативной базе ТЭК (ГОСТ, СТО, ФНП) с миллисекундным временем отклика через stdio / JSON-RPC."),
        ("DrillSense (Предиктивное бурение): ", "ML-модель прогнозирования скорости проходки ROP по телеметрии MWD/LWD скважины 15/9-F-9 A (Equinor Volve Data Village). Точность R2 = 0.892, MAE = 2.41 м/ч."),
        ("PumpGuard (IoT MLOps): ", "высоконагруженный конвейер мониторинга фонда 500 погружных насосов УЭЦН (ClickHouse 5000 строк/сек + Spark + MLflow), расчет риска отказа за 72 часа до инцидента."),
        ("B2B Contract Risk Checker: ", "экспресс-аудит коммерческих договоров с детекцией Red Flags стоп-факторов и автогенерацией официального протокола разногласий на Pydantic v2."),
        ("StockFlow (Ритейл): ", "системная инженерия ИИ-агента прогнозирования спроса на 14 дней вперед в команде из 7 человек, точность ~75%, целевое снижение дефицита полок на 15%."),
        ("Telecom Churn & Bank Marketing: ", "модель прогнозирования оттока клиентов (XGBoost, AUC-ROC = 0.8425) и интерактивный BI-дашборд в Yandex DataLens по 41 000+ контактов.")
    ]
    for r_t, r_d in rd_projects:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.left_indent = Inches(0.15)
        p_r.paragraph_format.space_after = Pt(1.5)
        p_r.paragraph_format.line_spacing = 1.12
        r_rt = p_r.add_run(f"– {r_t}")
        r_rt.bold = True
        r_rt.font.size = Pt(8.5)
        r_rd = p_r.add_run(r_d)
        r_rd.font.size = Pt(8.5)

    # Управленческий трек-рекорд
    add_h("УПРАВЛЕНЧЕСКИЙ ТРЕК-РЕКОРД И ПРЕДПРИНИМАТЕЛЬСТВО")
    mgmt_items = [
        "Организация масштабных офлайн-мероприятий (до 3 000 участников) с 10+ субподрядчиками: 100% соблюдение сметы, успешный кризис-менеджмент.",
        "Федеральная партнерская интеграция с ФК «Зенит»: прямые переговоры с PR-отделом клуба, федеральные охваты, сюжет на телеканале «78».",
        "Запуск товарного e-commerce бизнеса: CustDev ниши, аудит фабрик, брендинг, логистика, аналитика, рост продаж на 300%.",
        "Организация футбольного турнира полного цикла в СПб на 200+ участников (12 команд, спонсорские контракты на 500 тыс. руб., плановая прибыль)."
    ]
    for m in mgmt_items:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.left_indent = Inches(0.15)
        p_m.paragraph_format.space_after = Pt(1.5)
        p_m.paragraph_format.line_spacing = 1.1
        r_m = p_m.add_run(f"• {m}")
        r_m.font.size = Pt(8.5)

    # Образование
    add_h("ОБРАЗОВАНИЕ И КВАЛИФИКАЦИЯ")
    edu_items = [
        ("2025 — 2027 (Магистратура): Национальный исследовательский университет ИТМО, Санкт-Петербург", "Факультет технологий менеджмента и инноваций. Направление: «Технологии и стратегии бизнес-трансформации, Инноватика». Научный фокус: применение LLM и мультиагентных систем в корпоративной аналитике и трансформации."),
        ("2021 — 2025 (Бакалавриат): СПбГУТ им. проф. М.А. Бонч-Бруевича, Санкт-Петербург", "ЦЭУБИ. Направление: «Менеджмент цифровых технологий». ВКР: «Анализ потребительского поведения на основе больших данных»."),
        ("Профессиональная переподготовка: СПбГУТ им. Бонч-Бруевича", "Программа «Основы программирования и технологий искусственного интеллекта в экономике и финансах». Квалификация: Специалист по информационным системам.")
    ]
    for ed_t, ed_d in edu_items:
        p_ed = doc.add_paragraph()
        p_ed.paragraph_format.left_indent = Inches(0.15)
        p_ed.paragraph_format.space_after = Pt(1.5)
        p_ed.paragraph_format.line_spacing = 1.1
        r_edt = p_ed.add_run(f"• {ed_t}\n  ")
        r_edt.bold = True
        r_edt.font.size = Pt(8.5)
        r_edd = p_ed.add_run(ed_d)
        r_edd.font.size = Pt(8.5)

    # Достижения
    add_h("ДОСТИЖЕНИЯ И ДОПОЛНИТЕЛЬНАЯ ИНФОРМАЦИЯ")
    extra_items = [
        "Двукратный финалист Финансовой олимпиады МГУ.",
        "Выпускник Лидерской программы Evercode: Business Development Manager.",
        "Кандидат в мастера спорта (КМС) по дзюдо; член сборной Университета ИТМО по футболу.",
        "Языки: Русский (родной), Английский (B2 — Upper-Intermediate / профессиональная ИТ-терминология)."
    ]
    for ex in extra_items:
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.left_indent = Inches(0.15)
        p_ex.paragraph_format.space_after = Pt(1.5)
        p_ex.paragraph_format.line_spacing = 1.1
        r_ex = p_ex.add_run(f"✔ {ex}")
        r_ex.font.size = Pt(8.5)

    output_path = Path(__file__).resolve().parent / "Резюме_Лаврентий_Ямпуров_AI_PM.docx"
    doc.save(str(output_path))
    print(f"Резюме успешно сохранено: {output_path}")
    return str(output_path)


if __name__ == "__main__":
    create_resume()
