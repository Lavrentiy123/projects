import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Corporate Color Palette (Gazprom Neft inspired)
NAVY_BLUE = RGBColor(0, 51, 102)     # #003366
PRIMARY_BLUE = RGBColor(0, 121, 194) # #0079C2
DARK_TEXT = RGBColor(30, 41, 59)     # #1E293B
LIGHT_BG = RGBColor(244, 247, 251)   # #F4F7FB
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(235, 243, 250)    # #EBF3FA
BORDER_COLOR = RGBColor(200, 218, 235)
ACCENT_GREEN = RGBColor(16, 149, 106)
ACCENT_ORANGE = RGBColor(220, 100, 30)

def add_header(slide, title_text, category="ДСИиУР ПАО «ГАЗПРОМ НЕФТЬ» | АРХИТЕКТУРА ИИ-АГЕНТОВ"):
    # Header box
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_cat = tf.paragraphs[0]
    r_cat = p_cat.add_run()
    r_cat.text = category.upper()
    r_cat.font.size = Pt(10)
    r_cat.font.bold = True
    r_cat.font.color.rgb = PRIMARY_BLUE
    
    p_title = tf.add_paragraph()
    r_title = p_title.add_run()
    r_title.text = title_text
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY_BLUE

def add_card(slide, left, top, width, height, title, content_lines, border_color=PRIMARY_BLUE, bg_color=CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    p_t = tf.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = title
    r_t.font.bold = True
    r_t.font.size = Pt(14)
    r_t.font.color.rgb = NAVY_BLUE
    
    for line in content_lines:
        p_c = tf.add_paragraph()
        p_c.space_before = Pt(6)
        r_c = p_c.add_run()
        r_c.text = line
        r_c.font.size = Pt(11)
        r_c.font.color.rgb = DARK_TEXT

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10.0)
    prs.slide_height = Inches(5.625) # 16:9 widescreen
    blank_layout = prs.slide_layouts[6]

    # --- SLIDE 1: Title Slide ---
    s1 = prs.slides.add_slide(blank_layout)
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(5.625))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY_BLUE
    bg.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(8.0), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "ДЕПАРТАМЕНТ ПО СТРАТЕГИИ, ИННОВАЦИЯМ И УСТОЙЧИВОМУ РАЗВИТИЮ"
    r0.font.size = Pt(12)
    r0.font.bold = True
    r0.font.color.rgb = PRIMARY_BLUE
    
    p1 = tf.add_paragraph()
    p1.space_before = Pt(10)
    r1 = p1.add_run()
    r1.text = "Сквозная система ИИ-агентов для стратегической аналитики"
    r1.font.size = Pt(28)
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(14)
    r2 = p2.add_run()
    r2.text = "Концепция, архитектурный дизайн и дорожная карта (Аналитик → Стратег → Презентатор)\nРазработчик: Лаврентий | Горизонт: Сентябрь – Декабрь"
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(190, 215, 240)

    # --- SLIDE 2: Проблематика и Бизнес-цель ---
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Проблематика и целевые бизнес-эффекты")
    add_card(s2, Inches(0.8), Inches(1.6), Inches(4.0), Inches(3.4), 
             "Текущие вызовы (As Is)", 
             ["• До 70% времени аналитиков уходит на сбор новостей, вычитку PDF и ручную верстку.",
              "• Огромный объем открытых данных: сотни источников, регуляторные изменения, отчеты конкурентов.",
              "• «Разрыв коммуникации»: стратегия формулируется на верхнем уровне, а отделы не видят своих конкретных задач."],
             border_color=ACCENT_ORANGE)
    add_card(s2, Inches(5.2), Inches(1.6), Inches(4.0), Inches(3.4), 
             "Целевое состояние (To Be)", 
             ["• Автономный мониторинг: Агент 1 круглосуточно собирает дайджесты и сигналы рисков.",
              "• Быстрое сценарное планирование: Агент 2 генерирует инициативы и тестирует риски в цикле Reflexion.",
              "• Персонализация под ДО и отделы: Агент 3 собирает адресные презентации .pptx по корпоративному шаблону."],
             border_color=ACCENT_GREEN)

    # --- SLIDE 3: Общая сквозная архитектура ---
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Сквозной конвейер взаимодействия 3-х агентов")
    add_card(s3, Inches(0.6), Inches(1.6), Inches(2.7), Inches(3.4),
             "1. Агент-Аналитик",
             ["• Парсинг отчетов и новостей",
              "• Layout-Aware RAG (Docling)",
              "• Сигналы раннего предупреждения",
              "• Бенчмарк конкурентов",
              "\nВыход: AnalystOutput (JSON)"],
             border_color=PRIMARY_BLUE)
    add_card(s3, Inches(3.65), Inches(1.6), Inches(2.7), Inches(3.4),
             "2. Агент-Стратег",
             ["• Входной профиль стратегии",
              "• Фреймворки: SWOT / PESTEL",
              "• Модуль критики (Reflexion)",
              "• Верификация рисков и сроков",
              "\nВыход: StrategyRoadmap (JSON)"],
             border_color=PRIMARY_BLUE)
    add_card(s3, Inches(6.7), Inches(1.6), Inches(2.7), Inches(3.4),
             "3. Агент-Презентатор",
             ["• Профили отделов (YAML)",
              "• Селекция релевантных инициатив",
              "• Перевод на язык подразделения",
              "• Генерация .pptx по шаблону",
              "\nВыход: Адресные слайды"],
             border_color=PRIMARY_BLUE)

    # --- SLIDE 4: Агент 1 Детально ---
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Агент 1: Рыночный Аналитик (Data Ingestion & Signals)")
    add_card(s4, Inches(0.8), Inches(1.6), Inches(4.1), Inches(3.4),
             "Ключевой функционал",
             ["• Интеллектуальный парсинг: извлечение сложных таблиц из годовых отчетов ТЭК через Docling.",
              "• Гибридный поиск: объединение семантических векторов и точного совпадения ключевых слов (BM25).",
              "• Кросс-энкодерный реранкинг: отсечение информационного шума до генерации текста."],
             border_color=PRIMARY_BLUE)
    add_card(s4, Inches(5.1), Inches(1.6), Inches(4.1), Inches(3.4),
             "Выходные артефакты",
             ["1. Отраслевой дайджест: тренды ТЭК, цены, технологии.",
              "2. Аналитическая справка: глубокий синтез темы.",
              "3. Сигналы раннего предупреждения: критические шоки (санкции, дефицит флота, сбои поставок).",
              "4. Карта стратегий: матрица ходов конкурентов (Роснефть, Лукойл, Aramco)."],
             border_color=PRIMARY_BLUE)

    # --- SLIDE 5: Агент 2 Детально ---
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Агент 2: Корпоративный Стратег (Reflexion & Synthesis)")
    add_card(s5, Inches(0.8), Inches(1.6), Inches(4.1), Inches(3.4),
             "Методология и цикл критики",
             ["• Интеграция корпоративных целей: горизонт планирования (2026–2030) и риск-аппетит.",
              "• Контур Reflexion (Generator-Critic):",
              "  1) Генератор предлагает стратегическую инициативу.",
              "  2) Критик (Риск-офицер) проверяет регуляторику, бюджет и сроки.",
              "  3) Корректировка до оценки качества >= 8.0."],
             border_color=PRIMARY_BLUE)
    add_card(s5, Inches(5.1), Inches(1.6), Inches(4.1), Inches(3.4),
             "Результаты работы Стратега",
             ["• Дерево стратегических инициатив (STRAT-01..N).",
              "• Маппинг ответственных подразделений (ИТ, Логистика, Бурение, Капстрой).",
              "• Оценка ожидаемого эффекта и ключевых KPI.",
              "• Программа митигации макроэкономических рисков."],
             border_color=PRIMARY_BLUE)

    # --- SLIDE 6: Агент 3 Детально ---
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Агент 3: Адаптер под отделы и Генератор презентаций")
    add_card(s6, Inches(0.8), Inches(1.6), Inches(4.1), Inches(3.4),
             "Смысловая адаптация (YAML Profiles)",
             ["• Библиотека профилей: каждый отдел имеет профиль с фокусом, KPI и профессиональным сленгом.",
              "• Фильтрация: из 20 глобальных инициатив отбираются только 3–4 релевантных конкретному блоку.",
              "• Трансляция смыслов: верхнеуровневая цель «Декарбонизация» для Логистики превращается в «Оптимизация маршрутов танкеров и СПГ-топливо»."],
             border_color=PRIMARY_BLUE)
    add_card(s6, Inches(5.1), Inches(1.6), Inches(4.1), Inches(3.4),
             "Автоматическая верстка слайдов",
             ["• Принцип пирамиды Минто: ключевой инсайт в заголовке слайда, далее фактура и цифры.",
              "• Программная генерация в python-pptx: прямое заполнение официального мастер-шаблона компании.",
              "• 100% соблюдение корпоративного стиля: шрифты, логотипы, цвета без ручных правок."],
             border_color=PRIMARY_BLUE)

    # --- SLIDE 7: Технологический стек ---
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Технологический стек платформы")
    add_card(s7, Inches(0.8), Inches(1.6), Inches(2.7), Inches(3.4),
             "Оркестрация и Логика",
             ["• LangGraph: циклические графы, StateGraph, Checkpointing.",
              "• Pydantic v2 + Instructor: гарантированный валидный JSON.",
              "• LiteLLM: единая шина доступа к любым LLM."],
             border_color=PRIMARY_BLUE)
    add_card(s7, Inches(3.65), Inches(1.6), Inches(2.7), Inches(3.4),
             "Данные и RAG",
             ["• Docling / Marker: распознавание таблиц и структуры PDF.",
              "• Tavily API: чистый веб-поиск без шума и рекламы.",
              "• Qdrant / Chroma: векторная база в локальном Docker."],
             border_color=PRIMARY_BLUE)
    add_card(s7, Inches(6.5), Inches(1.6), Inches(2.7), Inches(3.4),
             "Презентации и UI",
             ["• python-pptx: векторная сборка слайдов по шаблону.",
              "• Marp: рендеринг Markdown в презентации.",
              "• Streamlit: интерактивный веб-дашборд для демо."],
             border_color=PRIMARY_BLUE)

    # --- SLIDE 8: Безопасность и закрытый контур ---
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Безопасность и архитектурный мост в On-Premise")
    add_card(s8, Inches(0.8), Inches(1.6), Inches(4.1), Inches(3.4),
             "Фаза 1: Внешний контур (PoC)",
             ["• Разработка на открытых источниках и синтетике.",
              "• Полная конфиденциальность: корпоративные секреты не попадают в облачные API.",
              "• Высокая скорость разработки: запуск без многомесячных согласований с ИБ."],
             border_color=ACCENT_GREEN)
    add_card(s8, Inches(5.1), Inches(1.6), Inches(4.1), Inches(3.4),
             "Фаза 2: Перенос в закрытый контур",
             ["• Model-Agnostic код: замена облачного эндпоинта на локальный vLLM/TGI в .env файле.",
              "• Локальные модели: Qwen 2.5 (32B/72B), DeepSeek-V3, Llama 3.3 на собственных GPU.",
              "• Полный суверенитет данных внутри корпоративного периметра компании."],
             border_color=NAVY_BLUE)

    # --- SLIDE 9: Метрики и качество ---
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Фреймворк метрик и контроль галлюцинаций")
    add_card(s9, Inches(0.8), Inches(1.6), Inches(4.1), Inches(3.4),
             "Метрики RAG (Ragas / DeepEval)",
             ["• Faithfulness (> 0.95): подтвержденность каждого тезиса справки цитатой из источника.",
              "• Answer Relevancy (> 0.90): точность ответа на стратегический вопрос.",
              "• Context Precision (> 0.85): чистота контекста без спама и шума."],
             border_color=PRIMARY_BLUE)
    add_card(s9, Inches(5.1), Inches(1.6), Inches(4.1), Inches(3.4),
             "Бизнес-метрики и тесты",
             ["• Actionability Score: конкретность предложенных шагов, сроков и KPI.",
              "• Adoption Rate: доля сгенерированных тезисов, принятых аналитиком без правок.",
              "• Пирамида тестов: Pydantic валидация -> Интеграция графа -> LLM-as-a-judge."],
             border_color=PRIMARY_BLUE)

    # --- SLIDE 10: Дорожная карта ---
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Дорожная карта реализации (Сентябрь – Декабрь)")
    add_card(s10, Inches(0.6), Inches(1.6), Inches(2.0), Inches(3.4),
             "Сентябрь",
             ["• Агент 1: Аналитик",
              "• Парсинг PDF (Docling)",
              "• RAG + Поиск",
              "• Детектор Early Warnings",
              "• Тесты на синтетике"],
             border_color=PRIMARY_BLUE)
    add_card(s10, Inches(2.9), Inches(1.6), Inches(2.0), Inches(3.4),
             "Октябрь",
             ["• Агент 2: Стратег",
              "• Контур Reflexion",
              "• Промпты SWOT/PESTEL",
              "• Генератор инициатив",
              "• Модуль критики рисков"],
             border_color=PRIMARY_BLUE)
    add_card(s10, Inches(5.2), Inches(1.6), Inches(2.0), Inches(3.4),
             "Ноябрь",
             ["• Агент 3: Презентатор",
              "• Реестр профилей отделов",
              "• Сборщик python-pptx",
              "• Наполнение шаблонов",
              "• Тесты адаптации"],
             border_color=PRIMARY_BLUE)
    add_card(s10, Inches(7.5), Inches(1.6), Inches(2.0), Inches(3.4),
             "Декабрь",
             ["• Финал и Демо",
              "• Streamlit UI",
              "• Методологическое описание для ДО",
              "• Подготовка к On-Prem",
              "• Защита перед C-level"],
             border_color=ACCENT_GREEN)

    # --- SLIDE 11: Чек-лист ресурсов ---
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Чек-лист вводных данных для старта")
    add_card(s11, Inches(0.8), Inches(1.6), Inches(8.4), Inches(3.4),
             "Что необходимо запросить у команды ДСИиУР на первой встрече:",
             ["1. 2–3 эталонных исторических аналитических отчета ДСИиУР («золотой стандарт» структуры и Tone-of-Voice).",
              "2. Официальный корпоративный .pptx шаблон с мастер-слайдами, цветовой палитрой и шрифтами «Газпром нефти».",
              "3. Список 2–3 пилотных подразделений (например: ИТ, Логистика, Бурение) для первичного тестирования Агента 3.",
              "4. Перечень приоритетных открытых отраслевых источников (сайты министерств, агентства, порталы).",
              "5. Лимиты на API-ключи (Claude 3.5 Sonnet, GPT-4o, Tavily) или тестовая виртуальная машина для стенда."],
             border_color=PRIMARY_BLUE)

    # --- SLIDE 12: Финал ---
    s12 = prs.slides.add_slide(blank_layout)
    bg12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(5.625))
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = NAVY_BLUE
    bg12.line.fill.background()

    tb12 = s12.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(8.0), Inches(2.6))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    
    p12_0 = tf12.paragraphs[0]
    r12_0 = p12_0.add_run()
    r12_0.text = "ГОТОВНОСТЬ К РЕАЛИЗАЦИИ И СТАРТУ"
    r12_0.font.size = Pt(12)
    r12_0.font.bold = True
    r12_0.font.color.rgb = PRIMARY_BLUE
    
    p12_1 = tf12.add_paragraph()
    p12_1.space_before = Pt(10)
    r12_1 = p12_1.add_run()
    r12_1.text = "Спасибо за внимание! Готов обсудить детали."
    r12_1.font.size = Pt(26)
    r12_1.font.bold = True
    r12_1.font.color.rgb = WHITE
    
    p12_2 = tf12.add_paragraph()
    p12_2.space_before = Pt(14)
    r12_2 = p12_2.add_run()
    r12_2.text = "Кандидат / Разработчик: Лаврентий\nКонтакты: Telegram | Телефон | Репозиторий GitHub"
    r12_2.font.size = Pt(13)
    r12_2.font.color.rgb = RGBColor(190, 215, 240)

    out_path = os.path.abspath("c:/Users/Lavrentiy/Desktop/газпром/Презентация_ИИ_Агенты_ГазпромНефть.pptx")
    prs.save(out_path)
    print(f"Presentation saved to: {out_path}")

if __name__ == "__main__":
    create_presentation()
